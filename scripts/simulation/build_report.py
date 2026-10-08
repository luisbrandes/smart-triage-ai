import json
import math
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
RESULTS_PATH = ROOT / "artifacts" / "simulation" / "observed-results.json"
OUTPUT_PATH = ROOT / "Relatorio_Analise_Simulacao_SmartTriageIA.docx"
CHART_DIR = ROOT / "artifacts" / "simulation" / "report-assets"

NAVY = "17324D"
BLUE = "286A8B"
LIGHT_BLUE = "EAF3F7"
LIGHT_GRAY = "F3F5F6"
BORDER = "D9D9D9"
TEXT = "1F2933"
MUTED = "52606D"
GREEN = "4E8A65"
YELLOW = "D59B2D"
RED = "B64B4B"


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def set_run_font(run, name="Arial", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_alt_text(inline_shape, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_paragraph(doc, text, bold_lead=None, keep=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        set_run_font(paragraph.add_run(item))


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=9, color=MUTED)
    run.italic = True


def add_table(doc, headers, rows, widths=None, font_size=8.5, header_fill=NAVY):
    compact = len(rows) > 25
    if compact:
        font_size = min(font_size, 6.8)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    if widths:
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
    set_table_borders(table)
    header = table.rows[0]
    mark_repeat_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = ""
        set_cell_shading(cell, header_fill)
        vertical_margin = 65 if compact else 100
        set_cell_margins(cell, vertical_margin, 85, vertical_margin, 85)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(value))
        set_run_font(run, size=font_size, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[index])
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.text = ""
            if row_index % 2:
                set_cell_shading(cell, LIGHT_BLUE)
            vertical_margin = 40 if compact else 80
            set_cell_margins(cell, vertical_margin, 75, vertical_margin, 75)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index != 1 and index != 2 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            set_run_font(run, size=font_size, color=TEXT)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def draw_priority_chart(counts, path):
    image = Image.new("RGB", (1400, 760), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "Distribuição das recomendações observadas", fill="#111111", font=font(38, True))
    values = [("Verde", counts["GREEN"], GREEN), ("Amarela", counts["YELLOW"], YELLOW), ("Vermelha", counts["RED"], RED)]
    max_value = max(value for _, value, _ in values)
    baseline = 650
    chart_top = 145
    for index, (label, value, color) in enumerate(values):
        x0 = 185 + index * 390
        x1 = x0 + 220
        height = int((baseline - chart_top) * value / max_value)
        draw.rounded_rectangle((x0, baseline - height, x1, baseline), radius=16, fill=f"#{color}")
        draw.text((x0 + 90, baseline - height - 55), str(value), fill="#111111", font=font(34, True), anchor="mm")
        draw.text((x0 + 110, baseline + 42), label, fill="#263238", font=font(28, True), anchor="mm")
        draw.text((x0 + 110, baseline + 78), f"{value / 50:.0%}", fill="#52606D", font=font(22), anchor="mm")
    draw.line((120, baseline, 1280, baseline), fill="#9AA5B1", width=3)
    image.save(path)


def draw_queue_chart(cases, path):
    image = Image.new("RGB", (1500, 850), "white")
    draw = ImageDraw.Draw(image)
    draw.text((65, 40), "Chegada e posição inicial na fila de revisão", fill="#111111", font=font(38, True))
    left, top, right, bottom = 120, 140, 1400, 730
    draw.rectangle((left, top, right, bottom), outline="#B8C2CC", width=2)
    colors = {"GREEN": GREEN, "YELLOW": YELLOW, "RED": RED}
    for rank in (1, 10, 20, 30, 40, 50):
        y = top + (rank - 1) / 49 * (bottom - top)
        draw.line((left, y, right, y), fill="#E4E7EB", width=2)
        draw.text((left - 18, y), str(rank), fill="#52606D", font=font(18), anchor="rm")
    for case in cases:
        arrival = case["arrivalOffsetMinutes"]
        rank = case["initialReviewQueuePosition"]
        x = left + arrival / 245 * (right - left)
        y = top + (rank - 1) / 49 * (bottom - top)
        color = colors[case["observedPriority"]]
        draw.ellipse((x - 8, y - 8, x + 8, y + 8), fill=f"#{color}", outline="#FFFFFF", width=2)
    draw.text((left + right >> 1, 790), "Minutos desde a primeira chegada", fill="#263238", font=font(24), anchor="mm")
    draw.text((35, (top + bottom) // 2), "Posição", fill="#263238", font=font(22, True), anchor="mm")
    for index, (label, color) in enumerate((("Vermelha", RED), ("Amarela", YELLOW), ("Verde", GREEN))):
        x = 965 + index * 165
        draw.ellipse((x, 72, x + 20, 92), fill=f"#{color}")
        draw.text((x + 28, 82), label, fill="#263238", font=font(17), anchor="lm")
    image.save(path)


def draw_means_chart(stats, path):
    image = Image.new("RGB", (1500, 820), "white")
    draw = ImageDraw.Draw(image)
    draw.text((65, 35), "Médias descritivas por recomendação", fill="#111111", font=font(38, True))
    priorities = [("GREEN", "Verde", GREEN), ("YELLOW", "Amarela", YELLOW), ("RED", "Vermelha", RED)]
    panels = [
        ("Frequência cardíaca média", "meanBpm", 130, "bpm"),
        ("SpO₂ média", "meanSpo2", 100, "%"),
        ("Temperatura média", "meanTemp", 41, "°C"),
    ]
    for panel_index, (title, key, maximum, unit) in enumerate(panels):
        x0 = 70 + panel_index * 480
        draw.text((x0, 125), title, fill="#263238", font=font(24, True))
        for index, (priority, label, color) in enumerate(priorities):
            y = 205 + index * 150
            value = stats[priority][key]
            width = int(330 * value / maximum)
            draw.rounded_rectangle((x0, y, x0 + width, y + 62), radius=12, fill=f"#{color}")
            draw.text((x0, y - 28), label, fill="#52606D", font=font(18, True))
            draw.text((x0 + width + 12, y + 31), f"{value:.1f} {unit}", fill="#111111", font=font(19, True), anchor="lm")
        draw.text((x0, 700), "Média do conjunto sintético", fill="#7B8794", font=font(17))
    image.save(path)


def age_on(birth_date, reference):
    born = datetime.fromisoformat(birth_date).replace(tzinfo=timezone.utc)
    age = reference.year - born.year
    if (reference.month, reference.day) < (born.month, born.day):
        age -= 1
    return age


def short_alerts(case):
    labels = []
    if case["alerts"]["suddenChange"] == "yes":
        labels.append("súbita")
    if case["alerts"]["breathing"] == "yes":
        labels.append("respiração")
    if case["alerts"]["consciousness"] == "yes":
        labels.append("consciência")
    return ", ".join(labels) if labels else "nenhum"


def status_pt(status):
    return {
        "AWAITING_NURSE_REVIEW": "Aguardando revisão",
        "CARE_WAITING": "Aguardando atendimento",
        "CALLED": "Chamado",
        "IN_SERVICE": "Em atendimento",
        "COMPLETED": "Concluído",
    }.get(status, status)


def priority_pt(priority):
    return {"GREEN": "Verde", "YELLOW": "Amarela", "RED": "Vermelha"}[priority]


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_border(title)

    for style_name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    first_section = doc.sections[0]
    header = first_section.header.paragraphs[0]
    clear_paragraph(header)
    header_run = header.add_run("SmartTriageIA  |  Avaliação experimental controlada")
    set_run_font(header_run, size=8.5, color=MUTED)
    footer = first_section.footer.paragraphs[0]
    clear_paragraph(footer)
    add_page_number(footer)


def add_landscape_section(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    return section


def add_portrait_section(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    return section


def build_report():
    data = json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    cases = data["cases"]
    patients = data["patients"]
    counts = Counter(case["observedPriority"] for case in cases)
    stats = {}
    for priority in ("GREEN", "YELLOW", "RED"):
        selected = [case for case in cases if case["observedPriority"] == priority]
        stats[priority] = {
            "count": len(selected),
            "percent": len(selected) / len(cases) * 100,
            "meanBpm": sum(case["vitals"]["bpm"] for case in selected) / len(selected),
            "meanSpo2": sum(case["vitals"]["spo2"] for case in selected) / len(selected),
            "meanTemp": sum(case["vitals"]["temperatura"] for case in selected) / len(selected),
        }
    agreements = sum(1 for case in cases if case["agreement"])
    reasons = Counter(
        reason["reasonCode"]
        for case in cases
        for run in case["analyses"]
        for reason in (run.get("recommendation") or {}).get("reasons", [])
    )
    reference_date = datetime.fromisoformat("2026-09-15T11:00:00+00:00")
    patient_by_key = {patient["key"]: patient for patient in patients}

    CHART_DIR.mkdir(parents=True, exist_ok=True)
    priority_chart = CHART_DIR / "priority-distribution.png"
    queue_chart = CHART_DIR / "queue-order.png"
    means_chart = CHART_DIR / "means-by-priority.png"
    draw_priority_chart(counts, priority_chart)
    draw_queue_chart(cases, queue_chart)
    draw_means_chart(stats, means_chart)

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_border(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(95)
    title.paragraph_format.space_after = Pt(18)
    title_run = title.add_run("Relatório de Análise da Simulação SmartTriageIA")
    set_run_font(title_run, size=28, bold=True, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    subtitle_run = subtitle.add_run("Avaliação experimental com 10 pacientes fictícios e 50 triagens")
    set_run_font(subtitle_run, size=15, color="000000")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(20)
    for line in (
        "Versão do experimento  SMARTTRIAGE EXPERIMENT 2026 01",
        "Conjunto de regras  V0 DEMO RULESET 002",
        "Execução observada em 8 de outubro de 2026",
    ):
        run = meta.add_run(line + "\n")
        set_run_font(run, size=10.5, color=MUTED)
    opening = doc.add_paragraph()
    opening.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    opening.paragraph_format.space_before = Pt(20)
    opening.paragraph_format.line_spacing = 1.15
    opening_run = opening.add_run(
        "Conclusão principal. O sistema reproduziu integralmente suas regras demonstrativas nos 50 casos e ordenou as duas filas por prioridade e tempo. O experimento também evidenciou limites importantes: temperatura, texto livre, idade e histórico não influenciam o motor; alterações extremas isoladas de FC ou SpO₂ alcançam no máximo amarelo; e não existe base clínica validada para interpretar a concordância observada como acurácia assistencial."
    )
    set_run_font(opening_run, size=11.5, bold=True)
    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.paragraph_format.space_before = Pt(30)
    notice_run = notice.add_run("Uso exclusivamente acadêmico e demonstrativo  Dados totalmente sintéticos")
    set_run_font(notice_run, size=10, bold=True, color=RED)
    doc.add_page_break()

    doc.add_heading("Sumário", level=1)
    toc_items = [
        "1 Introdução", "2 Metodologia", "3 Caracterização dos pacientes", "4 Caracterização dos 50 atendimentos",
        "5 Análise dos critérios de gravidade", "6 Análise da fila", "7 Análise do histórico dos pacientes",
        "8 Análise de consistência", "9 Casos de maior prioridade demonstrativa", "10 Falsos positivos e falsos negativos",
        "11 Matriz de resultados", "12 Análise estatística descritiva", "13 Segurança e limitações",
        "14 Pontos positivos observados", "15 Problemas e oportunidades de melhoria", "16 Conclusão",
        "17 Anexo de dados completos", "18 Anexo de evidências técnicas",
    ]
    add_bullets(doc, toc_items)

    doc.add_heading("1 Introdução", level=1)
    add_paragraph(doc, "O SmartTriageIA organiza identificação, consentimento, anamnese, coleta de sinais vitais, recomendação automática explicável e revisão profissional. A recomendação é preliminar e permanece separada da decisão da enfermagem. O sistema não implementa protocolo clínico validado, diagnóstico ou prescrição.")
    add_paragraph(doc, "O experimento foi concebido para observar o comportamento efetivo da implementação, incluindo persistência, regras, fila, estados, histórico e projeções usadas pelo frontend. A análise não procura provar que o produto funciona; procura registrar onde ele é previsível, onde é limitado e onde código e documentação divergem.")
    add_paragraph(doc, f"Foram processadas {len(cases)} triagens reais no ambiente de demonstração, associadas a {len(patients)} pessoas fictícias. A evidência foi coletada após a execução do motor, dos repositories e dos casos de uso que alimentam as páginas profissionais.")

    doc.add_heading("2 Metodologia", level=1)
    add_paragraph(doc, "O catálogo contém exatamente cinco atendimentos por paciente. As chegadas ocorreram em intervalos de cinco minutos, entre 08h00 e 12h05 no horário de Brasília, usando instantes UTC persistidos. Os cenários variaram queixa, intensidade, três alertas, FC, SpO₂ e temperatura.")
    add_paragraph(doc, "A prioridade esperada foi anotada antes da execução pela leitura independente das regras versionadas. Em seguida, cada caso percorreu sessão técnica de dispositivo, validação do payload, persistência de Reading, aceite explícito na triagem, criação de FC e SpO₂ clínicas, análise final, emissão de senha e consulta da projeção de enfermagem.")
    add_paragraph(doc, "Seis casos foram confirmados como decisões humanas separadas. A fila de atendimento chamou dois casos vermelhos; o primeiro foi iniciado e concluído, e o segundo permaneceu chamado. Nenhuma prioridade automática foi alterada para produzir um resultado desejado.")
    add_bullets(doc, [
        "Dados de identificação e documentos são marcadores sintéticos exclusivos.",
        "O motor executado foi runDeterministicFinalAnalysis com V0-DEMO-RULESET-002.",
        "A observação do frontend usou os mesmos DTOs e queries server-side consumidos pelas páginas da enfermagem.",
        "A concordância mede aderência à especificação demonstrativa, não validade clínica.",
    ])

    doc.add_heading("3 Caracterização dos pacientes", level=1)
    patient_rows = []
    for patient in patients:
        patient_rows.append([
            patient["key"], patient["displayName"], age_on(patient["birthDate"], reference_date), patient["sex"], 5,
        ])
    add_caption(doc, "Tabela 1  Perfis totalmente fictícios usados no experimento")
    add_table(doc, ["ID", "Nome fictício", "Idade", "Sexo", "Atendimentos"], patient_rows, [0.55, 2.7, 0.65, 1.0, 0.9], 9)
    add_paragraph(doc, "O fluxo aceita menores de idade porque a validação atual verifica apenas se a data é válida e não futura. Não existe modelagem de responsável ou consentimento específico para menores; por isso, o perfil adolescente é útil como evidência de uma limitação, não como autorização para uso real.")

    add_landscape_section(doc)
    doc.add_heading("4 Caracterização dos 50 atendimentos", level=1)
    add_paragraph(doc, "A tabela resume todas as entradas e resultados observados. Temperatura aparece porque foi recebida na Reading técnica; ela não foi convertida em VitalSignMeasurement nem considerada pelo motor.")
    case_rows = []
    for case in cases:
        case_rows.append([
            case["key"], patient_by_key[case["patientKey"]]["displayName"], case["mainComplaint"],
            case["vitals"]["bpm"], case["vitals"]["spo2"], f'{case["vitals"]["temperatura"]:.1f}',
            short_alerts(case), priority_pt(case["expectedPriority"]), priority_pt(case["observedPriority"]),
        ])
    add_caption(doc, "Tabela 2  Entradas e recomendações dos 50 atendimentos")
    add_table(doc, ["Caso", "Paciente", "Queixa", "FC", "SpO₂", "Temp", "Alertas", "Esperada", "Observada"], case_rows, [0.55, 1.55, 2.45, 0.45, 0.5, 0.5, 1.05, 0.75, 0.75], 7.5)

    add_portrait_section(doc)
    doc.add_heading("5 Análise dos critérios de gravidade", level=1)
    doc.add_heading("5.1 Critérios realmente consumidos", level=2)
    add_paragraph(doc, "O motor usa somente intensidade subjetiva, as três respostas de alerta, FC e SpO₂. O texto da queixa, descrição adicional, duração, idade, sexo, medicamentos, alergias, condições preexistentes, temperatura e histórico de triagens não compõem o NormalizedAnalysisInput.")
    rule_rows = [
        ["Alerta positivo", "Qualquer um dos três alertas", "Vermelha", "Prevalece sobre os demais fatores"],
        ["Intensidade alta", "7 ou mais sem alerta", "Amarela", "Adiciona razão de intensidade"],
        ["Intensidade intermediária", "4 a 6 sem alerta", "Amarela", "Não depende dos sinais"],
        ["Frequência cardíaca", "100 bpm ou mais", "Eleva verde para amarela", "Nunca eleva para vermelha"],
        ["SpO₂", "95 por cento ou menos", "Eleva verde para amarela", "Nunca eleva para vermelha"],
        ["Temperatura", "Sem regra", "Sem impacto", "Preservada apenas na Reading técnica"],
        ["Texto e histórico", "Sem regra", "Sem impacto", "Não chegam ao motor"],
    ]
    add_caption(doc, "Tabela 3  Critérios observados no código e confirmados pela simulação")
    add_table(doc, ["Critério", "Condição demonstrativa", "Efeito", "Limite observado"], rule_rows, [1.25, 1.75, 1.45, 2.1], 8.5)
    doc.add_heading("5.2 Alterações isoladas e combinadas", level=2)
    add_paragraph(doc, "SIM-042, com FC 220 e intensidade 2, recebeu amarelo. SIM-046, com SpO₂ 50 e intensidade 2, também recebeu amarelo. SIM-050 combinou FC 110 e SpO₂ 92 e permaneceu amarelo. Portanto, o motor acumula razões explicativas, mas não um escore que eleve múltiplas alterações a vermelho.")
    add_paragraph(doc, "SIM-041 registrou temperatura 40,0 °C e recebeu verde, pois temperatura não integra a análise clínica. Esse resultado é coerente com o código, mas evidencia uma diferença relevante entre o payload técnico disponível e os dados usados pela recomendação.")
    add_paragraph(doc, "SIM-043 informou intensidade 10 sem alerta positivo e recebeu amarelo. Vermelho depende exclusivamente de uma resposta positiva nas três perguntas demonstrativas, não da combinação de intensidade e sinais vitais.")
    doc.add_heading("5.3 Explicabilidade", level=2)
    add_paragraph(doc, f"Foram persistidas razões rastreáveis. Os gatilhos de FC apareceram {reasons['DEMO_HEART_RATE_ATTENTION']} vezes e os de SpO₂ {reasons['DEMO_SPO2_ATTENTION']} vezes. Os 15 casos vermelhos receberam a razão de coleta associada a alerta. O caso SIM-045 preservou cinco razões: dois alertas, solicitação de coleta, FC e SpO₂.")

    doc.add_heading("6 Análise da fila", level=1)
    add_paragraph(doc, "A fila de revisão não usa apenas ordem de chegada. Ela ordena primeiro Vermelha, depois Amarela e depois Verde; dentro de cada grupo, usa o instante de emissão da senha mais antigo. A simulação confirmou FIFO dentro das três prioridades.")
    shape = doc.add_picture(str(queue_chart), width=Inches(6.65))
    set_alt_text(shape, "Gráfico de dispersão com 50 chegadas e posições na fila, agrupadas em vermelho, amarelo e verde.")
    add_caption(doc, "Figura 1  A prioridade reorganiza a fila e o tempo desempata dentro de cada grupo")
    qa_rows = [
        ["O sistema usa apenas chegada?", "Não. A prioridade vem antes do tempo."],
        ["A gravidade demonstrativa influencia?", "Sim. A recomendação define o grupo da fila de revisão."],
        ["Um caso vermelho posterior ultrapassa um verde?", "Sim. SIM-003, às 08h10, ficou em 1º; SIM-001, às 08h00, ficou em 36º."],
        ["Como ocorre o desempate?", "Pela senha emitida mais cedo dentro da mesma prioridade."],
        ["O estado altera a fila?", "Sim. Seis confirmações retiraram casos da revisão e os inseriram no atendimento."],
        ["Existe risco de espera indefinida?", "Há risco teórico para prioridades menores sob chegada contínua de maiores; não há aging. Não foi demonstrado como espera infinita neste ensaio finito."],
    ]
    add_caption(doc, "Tabela 4  Respostas sobre o comportamento da fila")
    add_table(doc, ["Pergunta", "Resposta observada"], qa_rows, [2.2, 4.1], 9)
    add_paragraph(doc, "Na fila de atendimento, seis casos confirmados foram ordenados Vermelha, Amarela e Verde. A primeira chamada foi T-20260915-003 e a segunda T-20260915-006, exatamente os dois vermelhos por ordem de entrada. Ao final, havia 44 triagens em revisão, quatro aguardando atendimento, uma chamada e uma concluída.")

    doc.add_heading("7 Análise do histórico dos pacientes", level=1)
    add_paragraph(doc, "Cada Patient persistiu cinco Triage distintas, com anamnese, leitura, sessão, medições, análise, recomendação, senha e auditoria próprias. As relações permaneceram corretas e não foi encontrado cruzamento de sinais ou respostas entre atendimentos.")
    add_paragraph(doc, "O histórico anterior não influencia a recomendação atual. O motor recebe apenas o snapshot da anamnese mais recente da própria triagem e a sessão clínica aceita mais recente daquela triagem. Isso explica por que um mesmo paciente pode obter verde, amarelo ou vermelho exclusivamente conforme o caso atual.")
    add_paragraph(doc, "A interface profissional não apresenta uma visão longitudinal das cinco triagens do mesmo paciente. Além disso, o intake web atual cria um novo Patient em cada nova identificação e não busca uma pessoa existente. A associação múltipla foi preparada diretamente pelo experimento; portanto, persistência longitudinal existe no schema, mas o fluxo de usuário ainda não a constrói nem a exibe.")

    history_rows = []
    by_patient = defaultdict(list)
    for case in cases:
        by_patient[case["patientKey"]].append(case)
    for patient in patients:
        sequence = " → ".join(priority_pt(case["observedPriority"])[0] for case in by_patient[patient["key"]])
        history_rows.append([patient["key"], patient["displayName"], sequence, "Não influenciou"])
    add_caption(doc, "Tabela 5  Sequência de prioridades por pessoa  V verde  A amarela  R vermelha")
    add_table(doc, ["ID", "Paciente", "Cinco triagens", "Uso do histórico"], history_rows, [0.55, 2.6, 1.55, 1.4], 8.8)

    doc.add_heading("8 Análise de consistência", level=1)
    add_paragraph(doc, "O comportamento foi determinístico. Casos com a mesma estrutura de entrada produziram o mesmo nível, independentemente de paciente, idade, sexo, queixa textual ou atendimento anterior. Todos os 50 resultados coincidiram com a expectativa derivada da regra versionada.")
    consistency_rows = [
        ["Alertas semelhantes", "SIM-003, SIM-013, SIM-021, SIM-035", "Vermelha em todos"],
        ["FC isolada", "SIM-011, SIM-042, SIM-047", "Amarela em todos"],
        ["SpO₂ isolada", "SIM-017, SIM-025, SIM-046", "Amarela em todos"],
        ["Múltiplos gatilhos sem alerta", "SIM-020, SIM-044, SIM-050", "Amarela; mais razões, sem escalada"],
        ["Mesmo paciente em casos distintos", "Cinco casos por paciente", "Resultado depende somente da triagem atual"],
    ]
    add_caption(doc, "Tabela 6  Comparações de consistência")
    add_table(doc, ["Comparação", "Casos", "Resultado"], consistency_rows, [1.65, 2.25, 2.35], 8.8)

    doc.add_heading("9 Casos de maior prioridade demonstrativa", level=1)
    add_paragraph(doc, "Quinze casos receberam recomendação vermelha. A expressão maior prioridade demonstrativa descreve apenas a saída do ruleset; não é afirmação de emergência clínica.")
    critical_rows = []
    for case in [item for item in cases if item["observedPriority"] == "RED"]:
        critical_rows.append([case["key"], case["mainComplaint"], short_alerts(case), f'{case["vitals"]["bpm"]}/{case["vitals"]["spo2"]}', case["initialReviewQueuePosition"], "Concordante"])
    add_caption(doc, "Tabela 7  Casos com recomendação vermelha")
    add_table(doc, ["Caso", "Queixa", "Alertas", "FC SpO₂", "Posição", "Regra"], critical_rows, [0.55, 2.0, 1.25, 0.8, 0.7, 0.95], 8.2)
    add_paragraph(doc, "Todos os casos vermelhos ficaram nas posições 1 a 15 da fila inicial. SIM-045, apesar de chegar às 11h40 no horário local, ficou em 14º e ultrapassou todos os amarelos e verdes anteriores. A posição relativa entre os vermelhos continuou temporal.")

    doc.add_heading("10 Falsos positivos e falsos negativos", level=1)
    add_paragraph(doc, "Não identificado no conjunto de testes, quando a referência é estritamente a especificação das regras atuais: houve 50 concordâncias e nenhuma divergência.")
    add_paragraph(doc, "Não é tecnicamente válido calcular falsos positivos ou falsos negativos clínicos. O experimento não possui padrão-ouro profissional, protocolo validado, desfecho real ou revisão independente por especialistas. Classificar SIM-041, SIM-046 ou SIM-047 como erro clínico seria extrapolar a evidência disponível.")
    add_paragraph(doc, "Esses casos devem ser tratados como sinais de risco de produto: temperatura 40,0 °C resultou verde; SpO₂ 50 e FC 240 resultaram amarelo. Eles demonstram o comportamento do motor e justificam avaliação clínica futura, mas não autorizam uma conclusão de segurança ou eficácia.")

    doc.add_heading("11 Matriz de resultados", level=1)
    matrix_rows = [
        ["Esperada verde", 15, 0, 0, 15],
        ["Esperada amarela", 0, 20, 0, 20],
        ["Esperada vermelha", 0, 0, 15, 15],
        ["Total observado", 15, 20, 15, 50],
    ]
    add_caption(doc, "Tabela 8  Matriz de concordância com a especificação demonstrativa")
    add_table(doc, ["Referência", "Verde", "Amarela", "Vermelha", "Total"], matrix_rows, [1.65, 0.85, 0.9, 0.9, 0.7], 9)
    add_paragraph(doc, f"Concordância: {agreements}/50, ou {agreements / 50:.1%}. Divergências: {50 - agreements}. A métrica expressa apenas reprodutibilidade do ruleset; como expectativa e motor derivam da mesma especificação, ela não é uma estimativa independente de desempenho clínico.")

    doc.add_heading("12 Análise estatística descritiva", level=1)
    shape = doc.add_picture(str(priority_chart), width=Inches(6.45))
    set_alt_text(shape, "Gráfico de barras com 15 recomendações verdes, 20 amarelas e 15 vermelhas.")
    add_caption(doc, "Figura 2  Distribuição equilibrada para observar os três níveis")
    stat_rows = []
    for priority in ("GREEN", "YELLOW", "RED"):
        value = stats[priority]
        stat_rows.append([priority_pt(priority), value["count"], f'{value["percent"]:.0f}%', f'{value["meanBpm"]:.1f}', f'{value["meanSpo2"]:.1f}', f'{value["meanTemp"]:.2f}'])
    add_caption(doc, "Tabela 9  Estatísticas por recomendação")
    add_table(doc, ["Prioridade", "N", "%", "FC média", "SpO₂ média", "Temp média"], stat_rows, [1.25, 0.55, 0.65, 1.0, 1.05, 1.1], 9)
    shape = doc.add_picture(str(means_chart), width=Inches(6.65))
    set_alt_text(shape, "Três painéis com médias de frequência cardíaca, SpO₂ e temperatura por recomendação.")
    add_caption(doc, "Figura 3  Médias descritivas do conjunto sintético")
    add_paragraph(doc, "A média de FC foi maior e a média de SpO₂ menor no grupo amarelo porque o catálogo concentrou alterações isoladas e combinadas sem alertas nesse grupo. O grupo vermelho foi determinado pelos alertas, inclusive com sinais normais; por isso suas médias não são necessariamente mais alteradas. A temperatura média não explica prioridade, pois o motor não a utiliza.")

    doc.add_heading("13 Segurança e limitações", level=1)
    add_bullets(doc, [
        "A amostra é sintética, pequena, intencionalmente balanceada e não representa prevalência de serviço real.",
        "As regras são demonstrativas e não foram validadas contra protocolo ou decisão profissional independente.",
        "Os valores extremos foram aceitos pelos limites técnicos do payload; isso não equivale a validação clínica da medição.",
        "A automação observou a camada de aplicação usada pelo frontend, mas não constituiu teste E2E de navegador com interação humana.",
        "Os cinco atendimentos por pessoa exigiram bootstrap direto porque o fluxo web não reutiliza Patient existente.",
        "A fila foi observada em janela finita; starvation foi inferido do algoritmo de prioridade estrita, não produzido como espera infinita.",
        "Não foram avaliados desempenho sob carga, segurança ofensiva, disponibilidade de rede ou calibração de hardware.",
    ])
    add_paragraph(doc, "O SmartTriageIA deve ser apresentado como ferramenta experimental de apoio à organização e à revisão de enfermagem. A recomendação automática não substitui profissional, protocolo validado ou decisão clínica.")

    doc.add_heading("14 Pontos positivos observados", level=1)
    positive_rows = [
        ["Determinismo", "50 de 50 resultados reproduziram as regras versionadas."],
        ["Rastreabilidade", "Leitura técnica, aceite, medições clínicas, análise, razões, senha e auditoria ficaram separados."],
        ["Fila", "Prioridade e FIFO interno foram observados nas duas filas."],
        ["Histórico", "Cinco triagens por paciente permaneceram relacionadas sem cruzamento de dados."],
        ["Decisão humana", "Seis confirmações criaram NurseReview e finalHumanPriority sem apagar recomendações."],
        ["Idempotência", "Nova execução de simulation:load detectou o lote completo e não duplicou registros."],
    ]
    add_caption(doc, "Tabela 10  Comportamentos sustentados por evidência")
    add_table(doc, ["Aspecto", "Evidência observada"], positive_rows, [1.35, 4.95], 9)

    doc.add_heading("15 Problemas e oportunidades de melhoria", level=1)
    issue_rows = [
        ["Temperatura não clínica", "SIM-041 registrou 40,0 °C e permaneceu verde; o detalhe profissional não mostra temperatura.", "Definir decisão de domínio e validação antes de incorporar temperatura ao motor."],
        ["Teto amarelo para sinais", "FC 240, SpO₂ 50 e múltiplas alterações sem alerta alcançaram no máximo amarelo.", "Submeter limiares e combinações a governança e validação clínica; não ajustar por intuição técnica."],
        ["Texto, idade e histórico ignorados", "As entradas persistem, mas não compõem NormalizedAnalysisInput.", "Documentar com mais destaque ou criar evolução somente após requisitos clínicos formais."],
        ["Histórico não visível", "Detalhe profissional mostra uma triagem, sem lista longitudinal do paciente.", "Projetar visão mínima de histórico com controle de acesso e minimização."],
        ["Reuso de paciente ausente", "O intake cria novo Patient por triagem; o experimento precisou de bootstrap direto.", "Definir identificação/deduplicação segura antes de alegar histórico longitudinal."],
        ["Encaminhamento inalcançável", "Confirmar/reclassificar move direto a CARE_WAITING, mas referToCare exige CLASSIFIED; a UI não oferece encaminhar em CARE_WAITING.", "Alinhar máquina de estados, repository, UI, testes e ADR em uma decisão única."],
        ["Documentação divergente", "Documentos ainda descrevem análise preliminar e transição CLASSIFIED que o fluxo atual não executa.", "Atualizar documentação e critérios depois da decisão de produto."],
        ["Sem aging na fila", "Prioridade estrita sempre antecede tempo.", "Avaliar política formal de espera e escalonamento antes de uso operacional."],
    ]
    add_caption(doc, "Tabela 11  Problemas observados e sugestões separadas")
    add_table(doc, ["Problema observado", "Evidência", "Sugestão de melhoria"], issue_rows, [1.55, 2.35, 2.45], 7.9)

    conclusion_heading = doc.add_heading("16 Conclusão", level=1)
    conclusion_heading.paragraph_format.page_break_before = True
    add_paragraph(doc, "O sistema diferenciou os três níveis definidos por suas regras e manteve comportamento consistente em 50 execuções. Todos os alertas positivos geraram vermelho; intensidade e sinais técnicos produziram amarelo conforme os limiares; entradas sem gatilho produziram verde.")
    add_paragraph(doc, "Os casos de maior prioridade demonstrativa foram identificados e colocados antes dos demais. A ordem de chegada influenciou apenas dentro da mesma prioridade. Na fila de atendimento, a prioridade humana confirmada também precedeu o tempo.")
    add_paragraph(doc, "O histórico foi persistido corretamente por triagem e sem mistura, mas não influenciou a análise nem apareceu como visão longitudinal. O fluxo web tampouco reutiliza Patient existente, o que limita a aplicação prática desse histórico.")
    add_paragraph(doc, "A principal conclusão é dupla: a implementação é previsível em relação ao ruleset V0 e rastreia bem os registros, porém ainda é inadequado inferir segurança clínica. Temperatura ignorada, teto amarelo para alterações fisiológicas, ausência de contexto longitudinal e inconsistência do encaminhamento exigem decisão de produto, domínio e validação profissional antes de qualquer ampliação.")

    add_landscape_section(doc)
    doc.add_heading("17 Anexo de dados completos", level=1)
    add_paragraph(doc, "A tabela reúne as 50 triagens na ordem de chegada. Horários estão no fuso de Brasília. Posição refere-se à fila inicial de revisão antes das decisões humanas.")
    annex_rows = []
    for case in cases:
        arrival = datetime.fromisoformat(case["arrivalAt"].replace("Z", "+00:00"))
        local_hour = arrival.hour - 3
        annex_rows.append([
            case["key"], case["patientKey"], f"{local_hour:02d}:{arrival.minute:02d}", case["subjectiveIntensity"],
            short_alerts(case), case["vitals"]["bpm"], case["vitals"]["spo2"], f'{case["vitals"]["temperatura"]:.1f}',
            priority_pt(case["expectedPriority"]), priority_pt(case["observedPriority"]), case["initialReviewQueuePosition"], status_pt(case["persistedStatus"]),
        ])
    add_caption(doc, "Tabela 12  Dados integrais do conjunto experimental")
    add_table(doc, ["Caso", "Pessoa", "Chegada", "Int", "Alertas", "FC", "SpO₂", "Temp", "Esperada", "Observada", "Fila", "Estado final"], annex_rows, [0.55, 0.55, 0.65, 0.4, 1.05, 0.4, 0.45, 0.5, 0.72, 0.72, 0.45, 1.2], 7.2)

    add_portrait_section(doc)
    doc.add_heading("18 Anexo de evidências técnicas", level=1)
    doc.add_heading("18.1 Estrutura exercitada", level=2)
    technical_rows = [
        ["Dado técnico", "DeviceCollectionSession → Reading", "BPM, SpO₂ e temperatura"],
        ["Aceite clínico", "HardwareCollectionSession → VitalSignMeasurement", "FC e SpO₂ aceitas com origem física"],
        ["Análise", "AnalysisRun → AnalysisRecommendation → RecommendationReason", "Prioridade, ruleset e razões"],
        ["Fila de revisão", "QueueTicket e Triage AWAITING_NURSE_REVIEW", "Projeção por recomendação e tempo"],
        ["Decisão humana", "NurseReview e finalHumanPriority", "Registro separado da recomendação"],
        ["Atendimento", "CareAttendance", "WAITING, CALLED, IN_SERVICE e COMPLETED"],
    ]
    add_caption(doc, "Tabela 13  Cadeia de persistência observada")
    add_table(doc, ["Etapa", "Entidades", "Evidência"], technical_rows, [1.25, 2.8, 2.25], 8.5)

    doc.add_heading("18.2 Entradas HTTP relacionadas", level=2)
    add_paragraph(doc, "A aplicação expõe POST /api/triage/session/start, POST /api/readings, POST /api/triage/session/cancel, POST /api/triage/session/release e GET /api/triage/session/[sessionId]. O experimento automatizado não enviou HTTP: ele chamou os mesmos casos de uso e repositories server-side para eliminar dependência de rede e controlar o lote. A validação do payload de Reading foi executada antes da persistência.")
    add_paragraph(doc, "A projeção observada pelo frontend foi listNursingQueue, getNursingTriageDetail e listCareQueue, exatamente as consultas usadas nas páginas /enfermagem, /enfermagem/triagens/[triageId] e /enfermagem/atendimento.")

    doc.add_heading("18.3 Estados observados", level=2)
    state_rows = [
        ["Fila inicial", "50 AWAITING_NURSE_REVIEW"],
        ["Após seis confirmações", "44 AWAITING_NURSE_REVIEW e 6 CARE_WAITING"],
        ["Após primeira chamada e conclusão", "1 COMPLETED"],
        ["Após segunda chamada", "1 CALLED e 4 CARE_WAITING"],
        ["Leituras clínicas", "50 sessões PHYSICAL_DEVICE completas e 100 medições aceitas"],
        ["Sessões técnicas", "50 sessões concluídas e uma WAITING pronta para novo ciclo"],
    ]
    add_caption(doc, "Tabela 14  Estados capturados pelo experimento")
    add_table(doc, ["Momento", "Estado observado"], state_rows, [2.0, 4.3], 9)

    doc.add_heading("18.4 Artefatos reproduzíveis", level=2)
    add_bullets(doc, [
        "scripts/simulation/simulation-data.ts contém os 10 perfis e 50 cenários.",
        "npm run simulation:clean remove somente o lote marcado.",
        "npm run simulation:load executa o fluxo e é idempotente.",
        "artifacts/simulation/observed-results.json contém a evidência observada usada neste relatório.",
        "docs/SIMULATION_EXPERIMENT.md documenta segurança, execução e limites.",
    ])
    add_paragraph(doc, "Nenhum endpoint, log, relatório ou artefato contém credenciais. Todos os nomes, planos, documentos e situações são construções sintéticas para demonstração.")

    configure_document(doc)
    doc.core_properties.title = "Relatório de Análise da Simulação SmartTriageIA"
    doc.core_properties.subject = "Avaliação experimental com dados sintéticos"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = "Dados sintéticos sem validade clínica"
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_report()
