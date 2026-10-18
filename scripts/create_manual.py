from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('docs/apresentacao/Manual_de_Apresentacao_SmarTriageIA.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = '1F4E79'; LIGHT = 'DDEBF7'; PALE = 'F3F6F9'; BORDER = 'D9E2F3'; DARK = '17365D'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color=BORDER):
    tblPr = table._tbl.tblPr; b = tblPr.first_child_found_in('w:tblBorders')
    if b is None: b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = b.find(qn('w:'+edge))
        if el is None: el = OxmlElement('w:'+edge); b.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def set_cell(cell, text, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'; borders(t)
    trPr = t.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'), 'true'); trPr.append(hdr)
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i], h, True, 'FFFFFF', 9); shade(t.rows[0].cells[i], BLUE)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row): set_cell(cells[i], val, False, None, 8.5); shade(cells[i], 'FFFFFF' if ri%2==0 else PALE)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def bullet(doc, text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); return p

def numbered(doc, text):
    p=doc.add_paragraph(style='List Number'); p.add_run(text); return p

def code(doc, text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.space_after=Pt(5)
    r=p.add_run(text); r.font.name='Consolas'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('333333'); return p

def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)

def para(doc, text, bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.08
    if bold_lead and text.startswith(bold_lead): p.add_run(bold_lead).bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def page_break(doc): doc.add_page_break()

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10); styles['Normal'].paragraph_format.space_after=Pt(5)
for s in ['Title','Heading 1','Heading 2','Heading 3']:
    styles[s].font.name='Aptos Display'; styles[s].font.color.rgb=RGBColor(0,0,0)
styles['Title'].font.size=Pt(28); styles['Title'].font.bold=True
styles['Heading 1'].font.size=Pt(17); styles['Heading 1'].font.bold=True
styles['Heading 2'].font.size=Pt(13); styles['Heading 2'].font.bold=True
styles['Heading 3'].font.size=Pt(11); styles['Heading 3'].font.bold=True
for st in ['List Bullet','List Bullet 2','List Number']:
    styles[st].font.name='Aptos'; styles[st].font.size=Pt(10)

# header/footer
header=sec.header.paragraphs[0]; header.text='SmarTriageIA | Manual de apresentação'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor.from_string('667085')
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Material acadêmico e demonstrativo | Não substitui avaliação profissional'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('667085')

# cover
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(80)
r=p.add_run('SmarTriageIA'); r.font.name='Aptos Display'; r.font.size=Pt(34); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Sistema de Triagem Inteligente com Integração de Hardware e Apoio à Decisão Clínica'); r.font.size=Pt(15); r.bold=True
doc.add_paragraph('')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Manual de apresentação, estudo técnico e preparação para perguntas').font.size=Pt(12)
doc.add_paragraph('')
table(doc,['Identificação','Informação'],[['Autor','Luís Henrique Fagundes Brandes'],['Instituição','________________________________'],['Evento','________________________________'],['Data','____ / ____ / ______']], [1.5,4.7])
doc.add_paragraph('Versão de referência: estado do repositório auditado em 18/10/2026.').alignment=WD_ALIGN_PARAGRAPH.CENTER
page_break(doc)

heading(doc,'Como usar este manual',1)
para(doc,'Este manual foi escrito para preparar uma apresentação acadêmica ou tecnológica do SmarTriageIA. Ele combina explicação do produto, roteiro de demonstração, arquitetura, API, banco de dados, hardware, classificação, inteligência artificial, segurança e perguntas da banca. O código do repositório é a referência principal; quando uma informação não está confirmada nele, isso é indicado explicitamente.')
para(doc,'O sistema é experimental e demonstrativo. A classificação automática é uma recomendação preliminar, a sugestão da IA é opcional e a decisão final pertence à enfermagem. Não use este material para alegar validação clínica, certificação, conformidade regulatória ou autorização de uso assistencial.')
heading(doc,'Sumário',1); doc.add_paragraph('[[TOC]]'); page_break(doc)

heading(doc,'1 Visão geral do projeto',1)
para(doc,'O SmarTriageIA é uma aplicação web experimental para organizar a coleta inicial de informações em pronto atendimento, UPA ou hospital. O fluxo reúne identificação, consentimento, anamnese e sinais vitais; produz uma recomendação preliminar explicável; e entrega a triagem para revisão da enfermagem.')
heading(doc,'1.1 Problema e público',2)
para(doc,'O problema de produto é a coleta inicial fragmentada e a pouca rastreabilidade entre dados, recomendação automática e decisão profissional. O paciente usa um fluxo responsivo em tablet ou totem; a enfermagem utiliza uma fila e uma tela de revisão. O sistema não é prontuário, diagnóstico ou prescrição.')
heading(doc,'1.2 Objetivos e diferenciais',2)
for x in ['organizar a entrada do paciente em etapas claras;','preservar consentimento, respostas, medições, recomendações e decisões como histórico;','separar dado bruto do hardware de medição clínica aceita;','explicar a recomendação determinística com razões versionadas;','manter a revisão humana como etapa obrigatória;','permitir uma sugestão experimental do Groq sem alterar a decisão oficial.']: bullet(doc,x)
heading(doc,'1.3 Explicação de 30 segundos',2)
para(doc,'“O SmarTriageIA organiza a primeira etapa do atendimento: o paciente se identifica, aceita o termo, responde à anamnese e coleta sinais vitais. O sistema registra esses dados, aplica regras demonstrativas e apresenta uma prioridade preliminar explicável para a enfermagem. Um ESP32 pode enviar leituras físicas por HTTPS, mas o dado bruto só entra no fluxo clínico após aceitação explícita. A IA é experimental e serve como apoio; quem confirma a prioridade é o profissional.”')

heading(doc,'2 Roteiro de apresentação ao vivo',1)
para(doc,'Roteiro recomendado: 7 a 9 minutos. Ajuste a velocidade ao público e sempre anuncie quando a tela estiver em modo simulado ou quando uma leitura vier do ESP32 físico.')
table(doc,['Tempo','Mostrar','Fala sugerida','Destaque técnico / evitar'],[
['0:00–0:45','Capa e problema','“Em muitos fluxos, informação inicial, sinais vitais e decisão ficam separados. O projeto organiza essa passagem.”','Dizer que é protótipo experimental; não prometer diagnóstico.'],
['0:45–1:30','Arquitetura resumida','“Paciente e hardware chegam à mesma API; o PostgreSQL preserva o histórico.”','Next.js, Prisma e PostgreSQL são confirmados; não chamar de microsserviços.'],
['1:30–3:00','Paciente: identificação, consentimento e anamnese','“O fluxo exige consentimento e registra a triagem antes de avançar.”','Mostrar dados fictícios; não usar nome/documento real.'],
['3:00–4:15','Tela de coleta + ESP32','“A sessão é criada no site; o dispositivo descobre a sessão, envia leituras e o frontend consulta o estado.”','Explicar polling; não dizer que é WebSocket.'],
['4:15–5:00','Análise e fila','“A regra determinística gera recomendação e razões; a fila ordena o trabalho.”','Prioridade é preliminar e demonstrativa.'],
['5:00–6:30','Detalhe da enfermagem','“O enfermeiro pode revisar, confirmar ou reclassificar com justificativa.”','Decisão humana permanece separada.'],
['6:30–7:30','Sugestão Groq opcional','“A IA recebe um conjunto mínimo e devolve uma sugestão validada; falhas não bloqueiam o fluxo.”','Não dizer que a IA decide ou está validada clinicamente.'],
['7:30–8:30','Fechamento','“O próximo passo é validação controlada, não uso clínico imediato.”','Diferenciar realizado, experimental e planejado.']],[.8,1.5,2.8,1.6])
heading(doc,'2.1 Frases de segurança',2)
for x in ['“A classificação é uma recomendação preliminar; a decisão final pertence à enfermagem.”','“A sugestão do Groq é experimental, pode estar errada e não modifica a prioridade oficial.”','“O hardware físico foi integrado ao caminho técnico de recebimento; sensores e calibração precisam de validação própria.”','“Usamos somente dados fictícios nesta demonstração.”']: bullet(doc,x)

heading(doc,'3 Demonstração prática',1)
heading(doc,'3.1 Preparação',2)
for x in ['Abrir o site publicado e duas abas: paciente e enfermagem.','Confirmar banco e variáveis de ambiente no ambiente de deploy sem exibir valores.','Se for usar ESP32, confirmar Wi-Fi, HTTPS, certificado CA, deviceId e token fora do código.','Preparar um paciente fictício, uma queixa fictícia e uma triagem limpa.','Manter uma gravação ou modo simulado apenas como plano alternativo, identificando-o como simulado.']: bullet(doc,x)
heading(doc,'3.2 Passo a passo',2)
steps=['Paciente acessa o fluxo, aceita o termo e conclui identificação/anamnese.','Na coleta, o navegador chama POST /api/triage/session/start e guarda o sessionId.','O ESP32 consulta GET /api/device/session/current, recebe a sessão ACTIVE e envia uma ou mais leituras para POST /api/readings.','O frontend consulta GET /api/triage/session/{sessionId} a cada 2 segundos enquanto ACTIVE.','Após leituras válidas, o dispositivo envia POST /api/triage/session/release; a sessão física termina e o site atualiza a tela.','O paciente confirma a leitura aceita; o sistema cria a sessão clínica e medições aceitas conforme a regra do caso de uso.','A análise determinística gera prioridade preliminar e senha; a enfermagem revisa na fila.','Opcionalmente, o enfermeiro solicita a sugestão do Groq e compara com a recomendação, sem substituir a decisão.']
for s in steps: numbered(doc,s)
heading(doc,'3.3 Plano de contingência',2)
table(doc,['Falha','Como reconhecer','Resposta de apresentação'],[
['Wi-Fi/ESP32','dispositivo não aparece ou retorna 401/timeout','Explicar o fluxo e usar coleta simulada claramente identificada.'],
['API indisponível','tela de coleta mostra erro ou 5xx','Não insistir com dados reais; mostrar o contrato e o plano alternativo local.'],
['Sensor inválido','valores fora dos limites técnicos','Repetir a medição; se persistir, registrar como falha e usar modo manual/simulado.'],
['Banco indisponível','rotas retornam erro de persistência','Parar a demonstração persistente; não inventar confirmação.'],
['Groq indisponível','sugestão ausente/UNAVAILABLE','Continuar com recomendação determinística e revisão humana.'],
['Internet instável','polling falha ou páginas não carregam','Usar hotspot preparado e gravação; comunicar que é plano de contingência.']],[1.2,2.2,3.0])

heading(doc,'4 Arquitetura e tecnologias',1)
para(doc,'A aplicação é um monólito modular em Next.js App Router. Isso significa uma implantação única com fronteiras internas entre apresentação, aplicação, domínio, persistência e integrações; não há necessidade de alegar microsserviços.')
table(doc,['Tecnologia','Função no sistema','Motivo','Onde aparece'],[
['Next.js 16 / React 19','rotas web, Server Components e Route Handlers','unifica frontend e backend HTTP','src/app, src/features'],
['TypeScript 5.9','tipagem estática','reduz erros de contrato','src/, tsconfig.json'],
['Tailwind CSS 4','tokens e estilos da interface','consistência visual','src/app/globals.css'],
['Zod','validação de payloads','contratos explícitos na fronteira','src/server/application'],
['Prisma 7','mapeamento objeto-relacional e migrations','persistência tipada','prisma/, src/server/repositories'],
['PostgreSQL','fonte autoritativa dos dados','transações, relações e histórico','schema.prisma, migrations'],
['Vitest / ESLint / tsc','testes e gates de qualidade','verificação automatizada','package.json, *.test.ts'],
['Vercel','hospedagem documentada para a aplicação','deploy do Next.js','docs/VERCEL_REMOTE_HARDWARE_SETUP.md'],
['ESP32 + HTTP/HTTPS','coleta e envio físico','integração remota simples','contrato em docs/ESP32_API_CONTRACT.md'],
['Groq / llama-3.3-70b-versatile','sugestão experimental acionada pela enfermagem','apoio textual opcional','src/server/services/ai/groq-suggestion.ts']],[1.2,2.0,1.6,2.0])
heading(doc,'4.1 Papéis dos componentes',2)
table(doc,['Componente','Explicação simples'],[['Frontend','telas e interação do paciente/enfermagem; não acessa Prisma.'],['Backend/API','valida entrada, autoriza, orquestra casos de uso e devolve JSON.'],['Banco','guarda triagens, sessões, leituras, medições, análises, decisões e auditoria.'],['ESP32','microcontrolador que lê sensores e chama a API; não recebe dados do paciente.'],['Sensores','fontes físicas de sinais; o firmware normaliza valores antes do envio.'],['IA externa','serviço opcional que produz sugestão textual validada; não é o classificador oficial.']],[1.4,5.2])

heading(doc,'5 Fluxo completo dos dados',1)
para(doc,'Há dois caminhos que se encontram na API: o paciente conduz a triagem pelo navegador; o ESP32 envia uma leitura técnica associada a uma sessão de dispositivo. O backend autentica o dispositivo, valida os limites, grava Reading e só depois um caso de uso pode aceitar medições para o domínio clínico.')
heading(doc,'5.1 Diagrama de arquitetura',2)
table(doc,['Origem','Fronteira','Persistência / saída'],[
['Paciente','Next.js /paciente → Server Actions e rotas','Triage, Patient, ConsentRecord, Anamnesis'],
['ESP32','HTTPS + Bearer → /api/device/session/current e /api/readings','MedicalDevice, DeviceCollectionSession, Reading'],
['Aceite explícito','caso de uso vital-signs','HardwareCollectionSession + VitalSignMeasurement'],
['Análise','entrada normalizada → regras versionadas','AnalysisRun + AnalysisRecommendation'],
['IA opcional','ação profissional → Groq → validação Zod','AiSuggestion + AuditEvent'],
['Enfermagem','fila e detalhe server-side','NurseReview, QueueTicket, CareReferral']],[1.3,2.6,2.7])
code(doc,'Paciente → Frontend → API/Server Actions → PostgreSQL\nESP32 → Wi-Fi/HTTPS → API autenticada → Reading\nReading aceita explicitamente → VitalSignMeasurement → Análise → Fila → Enfermagem')
heading(doc,'5.2 Sequência física',2)
for x in ['O site inicia uma DeviceCollectionSession associada à triagem solicitante.','O ESP32 consulta a sessão ativa usando seu token, sem conhecer o triageId.','O dispositivo envia deviceId, sessionId e valores; o backend confere token, dispositivo, sessão, estado e limites.','O frontend lê uma projeção sanitizada da sessão; o navegador não acessa o banco.','Release encerra a sessão técnica. A aceitação clínica é outra etapa e preserva a distinção entre bruto e aceito.']: numbered(doc,x)

heading(doc,'6 Integração com o hardware',1)
para(doc,'O ESP32 é o ponto de aquisição e transporte. O repositório contém o contrato HTTP e a fronteira HardwareGateway, mas não contém o firmware completo do dispositivo nem um mapa de ligações elétricas. Assim, componentes, pinos, endereços I²C e versões de bibliotecas devem ser confirmados no projeto separado do hardware.')
heading(doc,'6.1 O que é confirmado e o que não é',2)
table(doc,['Item','Situação no repositório'],[['ESP32','Confirmado como identidade/protocolo do dispositivo no contrato HTTP.'],['BPM, SpO₂ e temperatura','Confirmados como campos recebidos por Reading.'],['MAX30102 / MLX90614','Referidos nos materiais do hardware, mas não confirmados no código deste repositório.'],['Pinos, endereços e LEDs','Não confirmados no repositório; verificar no firmware real.'],['Estados do firmware','Contrato recomenda IDLE/ACTIVE e não reutilizar sessão terminal; implementação física deve ser verificada no firmware.'],['HTTPS/CA','Obrigatório no contrato remoto; CA instalado e cadeia válida precisam ser testados no dispositivo.']],[1.5,5.1])
heading(doc,'6.2 Explicação para públicos diferentes',2)
para(doc,'Para o público leigo: o ESP32 funciona como um pequeno computador conectado aos sensores. Ele mede, confere se a leitura parece tecnicamente possível e envia um pacote JSON para o servidor. Para o público técnico: o firmware deve usar HTTPS, Bearer token individual, deviceId cadastrado, sessionId descoberto por polling e payload validado; nunca deve imprimir token ou dados sensíveis no monitor serial.')
heading(doc,'6.3 Investigação de falhas',2)
for x in ['401: token ausente, inválido, revogado ou dispositivo inativo.','403: token válido, mas deviceId do corpo não corresponde ao dispositivo autenticado.','400: JSON ou valores fora do contrato.','404: dispositivo/sessão não encontrada.','409: sessão incompatível, já encerrada ou sem estado permitido.','5xx: falha interna; registrar apenas código sanitizado e repetir com limite no firmware.']: bullet(doc,x)

heading(doc,'7 API e comunicação com o dispositivo',1)
para(doc,'API REST é um conjunto de rotas HTTP que recebe JSON e devolve JSON. POST cria ou altera; GET consulta. O token de dispositivo fica somente no firmware/ambiente seguro e nunca deve ser incluído neste manual.')
heading(doc,'7.1 Rotas do navegador paciente',2)
table(doc,['Método e rota','Payload','Sucesso','Erros importantes'],[
['POST /api/triage/session/start','{ "deviceId": "ESP32-001" }','201 { sessionId, status }','401 sessão paciente; 400 payload; 403 dispositivo não configurado; 409 sessão ativa; 404/500'],
['GET /api/triage/session/{sessionId}','sem corpo','200 com sessão, status, timestamps e readings[]','401 cookie paciente; 400 UUID; 404; 500'],
['POST /api/triage/session/cancel','{ "sessionId": "uuid" }','200 { sessionId, status: "CANCELLED" }','401; 400; 404; 500'],
['POST /api/triage/session/release','uso do dispositivo; body com deviceId/sessionId','200 { sessionId, status }','401; 400; 403; 404; 500']],[2.0,2.0,1.8,1.2])
heading(doc,'7.2 Rotas do firmware',2)
table(doc,['Método e rota','JSON enviado','Resposta / status'],[
['GET /api/device/session/current','nenhum','200 { status: "IDLE"|"ACTIVE", sessionId }; 401 sem autorização'],
['POST /api/readings','{ deviceId, sessionId, bpm, spo2, temperatura }','201 { success: true, readingId }; 400 inválido; 401 não autorizado; 403 mismatch; 404/409 sessão; 500'],
['POST /api/triage/session/release','{ deviceId, sessionId }','200 { sessionId, status }; 400/401/403/404/500']],[2.2,2.8,2.0])
code(doc,'Authorization: Bearer <TOKEN_FORA_DO_DOCUMENTO>\nAccept: application/json\nContent-Type: application/json\nUser-Agent: SmartTriageESP32/1.0')
para(doc,'Limites técnicos documentados: bpm inteiro entre 20 e 240; spo2 inteiro entre 50 e 100; temperatura finita entre 20 e 45. Esses limites não são uma validação clínica. O sessionId é obrigatório no fluxo remoto do firmware, embora o schema legado permita ausência em determinados cenários.')
heading(doc,'7.3 Exemplo didático sem credencial',2)
code(doc,'POST https://example.invalid/api/readings\nAuthorization: Bearer TOKEN_DE_EXEMPLO_NAO_UTILIZAVEL\n\n{"deviceId":"ESP32-001","sessionId":"550e8400-e29b-41d4-a716-446655440000","bpm":78,"spo2":98,"temperatura":36.7}\n\n201 {"success":true,"readingId":"550e8400-e29b-41d4-a716-446655440001"}')

heading(doc,'8 Banco de dados e modelagem',1)
para(doc,'O Prisma define relações e migrations para PostgreSQL. O modelo técnico registra o que chegou do dispositivo; o modelo clínico registra o que foi aceito no fluxo da triagem. Essa separação reduz o risco de tratar qualquer pacote bruto como medição clínica.')
table(doc,['Modelo','Responsabilidade','Relações principais'],[
['Patient','identidade mínima do paciente','triagens, consentimentos'],['Triage','agregado do atendimento inicial e máquina de estados','patient, anamnese, sessões, análises, fila, revisão'],['MedicalDevice','catálogo e autorização do dispositivo','capabilities, sessões técnicas, readings'],['DeviceCollectionSession','sessão técnica do totem','device, readings, triagem solicitante opcional'],['Reading','leitura bruta recebida','device e sessão técnica'],['HardwareCollectionSession','tentativa clínica de coleta','triage, device, sessão técnica fonte, medições'],['VitalSignMeasurement','medição aceita no domínio clínico','sessão clínica'],['AnalysisRun / Recommendation','execução e resultado determinístico','triage e razões'],['AiSuggestion','apoio externo experimental','triage, AnalysisRun e sessão clínica'],['NurseReview / QueueTicket / CareAttendance','decisão e continuidade do atendimento','triage']],[1.5,2.8,2.7])
heading(doc,'8.1 Diagrama entidade-relacionamento simplificado',2)
code(doc,'Patient 1 ─── N Triage 1 ─── N HardwareCollectionSession 1 ─── N VitalSignMeasurement\n                         │                         ▲\n                         │                         │ fonte aceita\nMedicalDevice 1 ─── N DeviceCollectionSession 1 ─── N Reading\n                         │\n                         └── solicitado por Triage (opcional)\nTriage 1 ─── N AnalysisRun 1 ─── 1 AnalysisRecommendation\nTriage 1 ─── N AiSuggestion / NurseReview / AuditEvent')

heading(doc,'9 Classificação e inteligência artificial',1)
heading(doc,'9.1 Regras determinísticas',2)
para(doc,'O motor demonstrativo recebe entrada normalizada e regras injetadas em src/demo/analysis-rules.ts. Ele produz uma prioridade RED, YELLOW ou GREEN, razões explicáveis, versão do conjunto de regras e fingerprint. A fila ordena RED, YELLOW, GREEN e depois a entrada mais antiga. A recomendação não é diagnóstico nem protocolo clínico validado.')
heading(doc,'9.2 Groq como apoio experimental',2)
para(doc,'A ação profissional da enfermagem pode solicitar uma sugestão ao endpoint compatível da Groq usando o modelo llama-3.3-70b-versatile. O serviço envia um conjunto mínimo: idade, sexo, queixa, intensidade/duração, alertas afirmativos, sinais aceitos e prioridade do sistema. O prompt exige linguagem de apoio, ausência de diagnóstico e declaração de dados insuficientes. A resposta é validada por Zod e persistida como AiSuggestion.')
table(doc,['Aspecto','Comportamento atual'],[['Acionamento','manual, pelo detalhe da triagem, não automático'],['Saída','prioridade sugerida RED/YELLOW/GREEN/BLUE, atenção HIGH/MODERATE/LOW, conclusão e fatores'],['Proteções','timeout de 6 s, JSON estruturado, validação de schema e limites de texto'],['Falha','registro UNAVAILABLE; recomendação determinística e revisão continuam'],['Decisão final','sempre NurseReview; a IA não altera a prioridade oficial'],['Ambiente','APP_MODE=demo; chave server-only GROQ_API_KEY; sem chave a sugestão fica indisponível']],[1.4,5.2])
heading(doc,'9.3 Conceitos para explicar',2)
for term,desc in [('IA generativa','modelo que produz texto novo a partir de uma entrada.'),('Modelo de linguagem','rede treinada para estimar e gerar sequências de linguagem.'),('Prompt','instruções e dados fornecidos ao modelo.'),('Inferência','execução do modelo para obter uma resposta.'),('Limitação','uma resposta plausível pode estar errada; por isso há validação, rastreabilidade e revisão humana.')]: para(doc,f'{term}: {desc}',bold_lead=term)

heading(doc,'10 Segurança, privacidade e confiabilidade',1)
for x in ['Autenticação de dispositivo por Bearer token individual, hash/pepper server-side e verificação de deviceId.','Cookies de paciente e enfermagem são separados, assinados, httpOnly e carregam apenas identificador opaco da triagem.','Entradas são validadas na fronteira; mensagens técnicas não vazam para a interface.','HTTPS é obrigatório no contrato remoto; validade do certificado e configuração real devem ser testadas por ambiente.','Prisma fica no servidor e o frontend não acessa banco.','Auditoria é append-only e registra ações relevantes com metadados sanitizados.','O manual e a demonstração devem usar dados fictícios; não inclua tokens, chaves, senhas, nomes ou documentos reais.','Consentimento é evidência de fluxo, não prova isolada de conformidade LGPD.']: bullet(doc,x)
heading(doc,'10.1 Confiabilidade',2)
para(doc,'Há idempotência e controle de estado para leituras/sessões em pontos relevantes, mas a confiabilidade física depende de Wi-Fi, firmware, sensores, calibração e operação. Em produção real ainda seriam necessários observabilidade, política de retry, gestão de dispositivos, testes de carga e validação de recuperação.')

heading(doc,'11 Limitações atuais e próximos passos',1)
table(doc,['Categoria','Itens'],[
['Confirmadas','produto experimental; sem validação clínica/certificação; firmware completo, pinos e sensores não estão neste repositório; IA externa opcional; temperatura técnica não é convertida automaticamente em medição clínica nesta aceitação atual.'],
['Não testado neste ambiente','suíte de integração PostgreSQL depende de TEST_DATABASE_URL; deploy real, limites da Groq, escala, recuperação de rede e calibração física não podem ser confirmados só pelo código local.'],
['Recomendadas','documentar firmware versionado; contrato de erro/retry; observabilidade; testes E2E; validação de sensores; revisão de segurança; gestão de expiração/rotação de tokens.'],
['Futuras','validação clínica formal; integração com identidade institucional; mais sensores; protocolo de atendimento; escalabilidade; modelo local ou provedor alternativo; interoperabilidade.']],[1.4,5.2])

heading(doc,'12 Perguntas do público e da banca',1)
qa=[
('Qual problema o projeto resolve?','Organiza a coleta inicial e melhora a rastreabilidade entre dados, recomendação e revisão. Não promete substituir o atendimento.'),
('Por que usar ESP32?','É um microcontrolador acessível e adequado para conectar sensores e rede; a escolha ainda precisa de validação de engenharia para uso contínuo.'),
('Como os sinais chegam ao sistema?','O ESP32 usa HTTPS e JSON, autentica-se com token e envia a leitura vinculada a uma sessão.'),
('Como garantem a sessão correta?','O backend confere token, deviceId, sessionId, vínculo do dispositivo e estado ACTIVE.'),
('O que ocorre se o Wi-Fi cair?','A leitura não é confirmada; o firmware deve tratar timeout e retry limitado. O fluxo pode recorrer a modo alternativo.'),
('E se o sensor estiver errado?','O servidor rejeita valores fora dos limites técnicos; isso não substitui calibração nem validação clínica.'),
('A IA toma a decisão final?','Não. Ela sugere; a recomendação oficial e a revisão da enfermagem permanecem separadas.'),
('Por que Groq?','Foi escolhida como integração externa experimental por API; o provedor não é garantia de disponibilidade ou correção.'),
('Por que não modelo local?','É uma alternativa futura; exige avaliar hardware, latência, manutenção e desempenho.'),
('Como funciona sem IA?','A regra determinística e a revisão humana continuam funcionando.'),
('O sistema substitui enfermeiro?','Não. A decisão final é profissional.'),
('Como validar a classificação?','Hoje ela é demonstrativa e versionada; validação clínica formal ainda é necessária.'),
('Os dados são reais ou simulados?','A demonstração deve usar dados fictícios; uma leitura física só deve ser chamada de física quando realmente vier do hardware.'),
('O hardware já foi testado?','O contrato e o caminho de recebimento estão implementados; sensores, firmware e ambiente devem ser comprovados em teste separado.'),
('Quais são os principais riscos?','Leitura inadequada, falha de rede, segredo exposto, dependência externa e interpretação indevida da recomendação.'),
('Como proteger dados?','Minimização, cookies opacos, separação de papéis, validação, HTTPS e logs sanitizados.'),
('Qual o custo?','Depende de hospedagem, banco, internet, hardware e provedor de IA; o repositório não fixa um orçamento de produção.'),
('Pode escalar para outras unidades?','A arquitetura modular ajuda, mas identidade, operação, segurança e validação precisam ser projetadas.'),
('O que falta para uso real?','Validação clínica, segurança e governança, integração institucional, operação monitorada e aprovação aplicável.'),
('Qual o diferencial?','A separação rastreável entre coleta técnica, aceite clínico, recomendação, IA e decisão humana.'),
('Por que PostgreSQL?','Oferece relações, transações e histórico adequados ao modelo; é acessado por Prisma.'),
('O frontend acessa o banco?','Não. Ele chama rotas e Server Actions.'),
('O que é Reading?','Dado bruto recebido do dispositivo, ainda não uma medição clínica aceita.'),
('O que é VitalSignMeasurement?','Registro clínico aceito por caso de uso, com origem, status técnico e sessão.'),
('O que acontece com sessão cancelada?','Novas leituras vinculadas são rejeitadas; o histórico permanece.'),
('Por que o token não aparece no site?','O token é credencial do dispositivo e deve permanecer fora do frontend e do repositório.'),
('O que significa 401?','Falha de autenticação, como token ausente ou inválido.'),
('O que significa 409?','Conflito de estado, como sessão já ativa ou leitura em sessão encerrada.'),
('Como a fila é ordenada?','RED, YELLOW, GREEN e depois o item mais antigo, como regra demonstrativa.'),
('A temperatura entra na análise?','A coleta técnica recebe temperatura; o aceite físico atual cria FC e SpO₂ aceitas, enquanto temperatura permanece técnica conforme o repositório.'),
('Como apresentar uma falha da Groq?','“O recurso é opcional; quando indisponível, o fluxo determinístico e humano continua.”'),
('O projeto está em conformidade LGPD?','Não se deve afirmar isso apenas por haver consentimento; é necessária análise jurídica e operacional.'),
('Qual o próximo experimento?','Repetir o fluxo com hardware versionado, dados fictícios e métricas técnicas antes de qualquer hipótese clínica.')]
for i,(q,a) in enumerate(qa,1):
    heading(doc,f'{i}. {q}',2); para(doc,a)

heading(doc,'13 Glossário técnico',1)
terms=[('API','Interface que permite que sistemas troquem dados.'),('Endpoint','URL específica de uma API.'),('HTTP/HTTPS','Protocolos de comunicação web; HTTPS adiciona criptografia em trânsito.'),('JSON','Formato textual estruturado usado nos payloads.'),('REST','Estilo de API baseado em recursos e métodos HTTP.'),('Frontend','Parte visual usada no navegador.'),('Backend','Parte server-side que aplica regras e acessa serviços.'),('PostgreSQL','Banco relacional usado como fonte autoritativa.'),('ORM','Camada que mapeia objetos para tabelas; Prisma é o ORM deste projeto.'),('ESP32','Microcontrolador com conectividade usado pelo hardware.'),('Sensor','Componente que mede uma grandeza física.'),('BPM','Batimentos por minuto.'),('SpO₂','Saturação periférica de oxigênio; no sistema é um valor recebido e tecnicamente validado.'),('I²C','Barramento comum de comunicação entre microcontrolador e sensores.'),('Token','Credencial que autoriza o dispositivo.'),('Sessão','Registro que delimita uma coleta e seu estado.'),('Autenticação','Verificação de quem faz a requisição.'),('Autorização','Verificação do que a requisição pode fazer.'),('Deploy','Publicação da aplicação em um ambiente executável.'),('Variável de ambiente','Configuração fornecida fora do código, como URL do banco.'),('Prompt','Instrução enviada ao modelo de linguagem.'),('Inferência','Execução do modelo para obter uma saída.'),('Classificação de risco','Prioridade preliminar produzida por regras demonstrativas; não é diagnóstico.')]
table(doc,['Termo','Definição'],terms,[1.6,5.0])

heading(doc,'14 Checklist final para o evento',1)
checks=['Sistema acessível e fluxo paciente/enfermagem verificado.','Banco disponível e migration aplicada no ambiente correto.','Variáveis configuradas sem expor valores.','API acessível pelo dispositivo e HTTPS validado.','ESP32 e sensores testados; firmware e CA confirmados.','Token individual configurado fora de logs e repositório.','Wi-Fi principal e hotspot alternativo prontos.','Paciente, queixa e leituras de demonstração são fictícios.','Fluxo start → readings → release testado.','Plano para cancelamento, erro de API e ausência de IA pronto.','Navegador, abas e telas abertas antes do público chegar.','Mensagem preparada para diferenciar físico de simulado.','Não capturar ou mostrar nomes, documentos, tokens ou chaves reais.','Dizer explicitamente que o sistema é experimental e não substitui o profissional.','Levar este manual e uma versão curta do contrato HTTP.']
for x in checks: bullet(doc,'☐ '+x)

heading(doc,'Referências do repositório',1)
para(doc,'As afirmações técnicas deste manual podem ser conferidas diretamente nos caminhos abaixo. O código atual prevalece sobre qualquer descrição anterior.')
table(doc,['Área','Arquivos de referência'],[
['Escopo e arquitetura','README.md; docs/MVP_SCOPE.md; docs/ARCHITECTURE.md; docs/ENGINEERING_RULES.md'],
['Domínio e fluxo','docs/DOMAIN_MODEL.md; docs/USER_FLOWS.md; docs/TRIAGE_STATE_MACHINE.md'],
['Hardware/API','docs/ESP32_API_CONTRACT.md; docs/VERCEL_REMOTE_HARDWARE_SETUP.md; src/app/api; src/server/application/device-collection-sessions; src/server/application/device-readings'],
['Persistência','prisma/schema.prisma; prisma/migrations; src/server/repositories'],
['Análise e IA','docs/ANALYSIS_ENGINE.md; docs/PHYSICAL_RECOLLECTION_AND_AI.md; src/domain/analysis; src/server/services/ai; src/server/repositories/prisma-ai-suggestion-repository.ts'],
['Segurança','docs/SECURITY_AND_PRIVACY.md; src/server/device-authentication.ts; src/server/security'],
['Qualidade','package.json; arquivos *.test.ts; vitest.integration.config.ts']],[1.5,5.1])
para(doc,'Validações locais registradas no estado auditado: Prisma validate, Prisma generate, testes unitários Vitest (110 testes), ESLint, TypeScript e build passaram. A suíte de integração depende de TEST_DATABASE_URL e não foi considerada executada sem esse banco de teste configurado. O deploy remoto e a calibração física devem ser verificados no ambiente real.')

# page number field
for section in doc.sections:
    p=section.footer.paragraphs[0]; p.add_run('  |  ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)

doc.core_properties.title='Manual de Apresentação SmarTriageIA'
doc.core_properties.subject='Manual técnico e roteiro de apresentação'
doc.core_properties.author='Luís Henrique Fagundes Brandes'
doc.core_properties.keywords='SmarTriageIA, triagem, hardware, ESP32, API, Prisma, Groq'
doc.save(OUT)
print(OUT)
